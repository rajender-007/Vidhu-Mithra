from __future__ import annotations

import re
from pathlib import Path
from typing import Any
from uuid import uuid4

from docx import Document


DATE_PATTERNS = [
    re.compile(r"\b(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})\b"),
    re.compile(r"\b(\d{1,2})\s+(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\s+(\d{4})\b", re.I),
]


def extract_text(path: Path, mime: str) -> tuple[str, int, list[str]]:
    if mime == "application/pdf" or path.suffix.lower() == ".pdf":
        import fitz
        pages = []
        with fitz.open(path) as pdf:
            for page in pdf:
                pages.append(page.get_text("text"))
        return "\n\n".join(pages), len(pages), [] if pages else ["No text was extracted; OCR may be required."]
    if path.suffix.lower() == ".docx" or "wordprocessingml" in mime:
        doc = Document(path)
        text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
        return text, 1, [] if text.strip() else ["No text was extracted."]
    if path.suffix.lower() in {".txt", ".csv", ".eml", ".json"}:
        return path.read_text(encoding="utf-8", errors="replace"), 1, []
    return "", 1, ["This file type needs OCR or a supported parser before analysis."]


def segment_clauses(text: str) -> list[dict[str, Any]]:
    pieces = [piece.strip() for piece in re.split(r"\n{2,}|(?=\b\d{1,2}[.)]\s)", text) if piece.strip()]
    clauses = []
    for index, piece in enumerate(pieces[:80], start=1):
        lower = piece.lower()
        if len(piece) < 25 and index > 1:
            continue
        high_terms = ("liable", "indemn", "penalty", "terminate", "forfeit", "exclusive", "unlimited", "any and all")
        medium_terms = ("notice", "renew", "confidential", "jurisdiction", "dispute", "payment")
        if any(term in lower for term in high_terms):
            color, level, category = "red", "HIGH", "risk"
        elif any(term in lower for term in medium_terms):
            color, level, category = "orange", "MEDIUM", "attention"
        else:
            color, level, category = "blue", "LOW", "reference"
        clauses.append({"id": str(uuid4()), "clause_no": str(index), "heading": f"Clause {index}", "text": piece[:5000], "category": category, "page_no": 1, "risk_color": color, "risk_level": level, "bbox": None})
    return clauses


def extract_entities(text: str) -> list[dict[str, Any]]:
    entities = []
    for match in re.finditer(r"(?:₹|Rs\.?\s*)([\d,]+(?:\.\d+)?)", text, re.I):
        entities.append({"type": "amount", "value": match.group(0), "confidence": 0.85})
    for pattern in DATE_PATTERNS:
        for match in pattern.finditer(text):
            entities.append({"type": "date", "value": match.group(0), "confidence": 0.8})
    return entities[:100]


def build_comments(clauses: list[dict[str, Any]]) -> list[dict[str, Any]]:
    comments = []
    for clause in clauses:
        if clause["risk_color"] in {"red", "orange"}:
            comments.append({
                "id": str(uuid4()), "clause_id": clause["id"], "color": clause["risk_color"], "risk_level": clause["risk_level"],
                "body": "This wording may deserve closer review in the context of the full agreement.",
                "why_it_matters": "It may allocate a broad obligation, remedy or timing requirement.",
                "checks": ["Check definitions and exceptions", "Check any cap, notice period or dispute process", "Compare this wording with related clauses"],
                "lawyer_question": "How does this clause apply to my facts, and can its scope be clarified or limited?", "status": "active",
            })
    return comments
